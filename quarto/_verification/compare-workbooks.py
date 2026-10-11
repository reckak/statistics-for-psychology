"""Compare a pre-change copy with current workbooks (read-only, openpyxl).

Usage: python quarto/_verification/compare-workbooks.py tmp/excel-migration/before
Checks every populated cell, cached result, tab order, merged range, validation
and chart's live data references. Styling is reviewed in native Excel previews.
"""
from pathlib import Path
import math
import re
import sys
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
baseline = Path(sys.argv[1]).resolve()

def formula(value):
    # Excel may add its compatibility prefix or quote simple sheet names.
    return re.sub(r"'([A-Za-z_][A-Za-z_0-9]*)'!", r'\1!', value.replace('_xlfn.', ''))

def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-12)
    return a == b

def bindings(chart):
    refs = []
    for series in chart.series:
        for key in ('cat', 'val', 'xVal', 'yVal'):
            data = getattr(series, key, None)
            if data:
                for ref_key in ('numRef', 'strRef'):
                    ref = getattr(data, ref_key, None)
                    if ref:
                        refs.append(formula(ref.f))
    return refs

errors = []
for n in range(1, 6):
    name = f'kapitola_{n:02}.xlsx'
    old = openpyxl.load_workbook(baseline/name)
    new = openpyxl.load_workbook(ROOT/'data'/name)
    old_values = openpyxl.load_workbook(baseline/name, data_only=True)
    new_values = openpyxl.load_workbook(ROOT/'data'/name, data_only=True)
    if old.sheetnames != new.sheetnames:
        errors.append((n, 'sheet names/order', old.sheetnames, new.sheetnames))
    count = 0
    for a in old:
        b = new[a.title]
        coords = {c.coordinate for sh in (a,b) for row in sh for c in row if c.value is not None}
        for addr in sorted(coords):
            x,y = a[addr],b[addr]
            vx,vy = x.value,y.value
            if x.data_type == 'f' and y.data_type == 'f':
                vx,vy = formula(vx),formula(vy)
                ca,cb = old_values[a.title][addr].value,new_values[a.title][addr].value
                if not same(ca, cb): errors.append((n,a.title,addr,'cached',ca,cb))
            if x.data_type != y.data_type or not same(vx,vy):
                errors.append((n,a.title,addr,'content/type',vx,vy))
            if x.number_format != y.number_format:
                errors.append((n,a.title,addr,'number format',x.number_format,y.number_format))
            count += 1
        if set(map(str,a.merged_cells.ranges)) != set(map(str,b.merged_cells.ranges)):
            errors.append((n,a.title,'merged cells'))
        def validations(sh):
            return sorted((str(v.sqref),v.type,v.operator,v.formula1,v.formula2,v.allowBlank,v.showErrorMessage)
                          for v in sh.data_validations.dataValidation)
        if validations(a) != validations(b): errors.append((n,a.title,'validation',validations(a),validations(b)))
        if len(a._charts) != len(b._charts): errors.append((n,a.title,'chart count'))
        for ca,cb in zip(a._charts,b._charts):
            if bindings(ca) != bindings(cb): errors.append((n,a.title,'chart bindings',bindings(ca),bindings(cb)))
            for axis in ('x_axis','y_axis'):
                aa,bb=getattr(ca,axis),getattr(cb,axis)
                if aa.numFmt != bb.numFmt: errors.append((n,a.title,axis,'axis format',aa.numFmt,bb.numFmt))
                if (aa.scaling.min,aa.scaling.max) != (bb.scaling.min,bb.scaling.max):
                    errors.append((n,a.title,axis,'axis limits'))
        if a.freeze_panes != b.freeze_panes: errors.append((n,a.title,'freeze panes',a.freeze_panes,b.freeze_panes))
    print(f'Chapter {n}: compared {count} populated cells, {len(old.sheetnames)} sheets')
for error in errors: print(error)
if errors: raise SystemExit(f'{len(errors)} differences require review')
print('PASS: content, formula results, sheet structure, validation and chart references preserved.')
